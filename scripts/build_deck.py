#!/usr/bin/env python3
"""Annotate walkthrough frames (yellow highlight) and assemble the 16-slide
PowerPoint at walkthrough/RADAR_Walkthrough.pptx."""
import json, os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FRAMES = os.path.join(ROOT, 'walkthrough', 'frames')
ANNOT = os.path.join(ROOT, 'walkthrough', 'annotated')
OUT = os.path.join(ROOT, 'walkthrough', 'RADAR_Walkthrough.pptx')
GOLD = (251, 191, 36)          # #FBBF24
GOLD_HEX = RGBColor(0xFB, 0xBF, 0x24)
INK = RGBColor(0x1F, 0x24, 0x30)
GREY = RGBColor(0x6B, 0x72, 0x80)

os.makedirs(ANNOT, exist_ok=True)
boxes = json.load(open(os.path.join(FRAMES, 'boxes.json')))

# ---------------------------------------------------------------------------
# Step 2 — annotate: one 4px rounded yellow rectangle per frame, nothing else
# ---------------------------------------------------------------------------
PAD = 8
for name, box in boxes.items():
    src = os.path.join(FRAMES, name)
    img = Image.open(src).convert('RGB')
    if box:
        d = ImageDraw.Draw(img)
        x0 = max(2, box['x'] - PAD); y0 = max(2, box['y'] - PAD)
        x1 = min(img.width - 2, box['x'] + box['width'] + PAD)
        y1 = min(img.height - 2, box['y'] + box['height'] + PAD)
        d.rounded_rectangle([x0, y0, x1, y1], radius=10, outline=GOLD, width=4)
    img.save(os.path.join(ANNOT, name))
print(f'annotated {len(boxes)} frames -> walkthrough/annotated/')

# ---------------------------------------------------------------------------
# Step 3 — assemble deck (16:9)
# ---------------------------------------------------------------------------
CAPTIONS = {
 '01_homepage.png': (
   'The front door — New Chat lands on a single question box over the whole supply base.',
   'One entry point: ask, and the workspaces answer.',
   'No hunting through menus — ask or jump straight to the worklist.'),
 '02_network_landing.png': (
   'Landing — the full network exists but nothing is prioritized yet.',
   'this is the before state — visibility without triage.',
   '500+ suppliers, no way to know where to start.'),
 '03_network_analyzed.png': (
   'ANALYZE surfaces only the suppliers that matter, with the selection rule stated on screen.',
   'one defensible answer instead of a wall — the "why these" is written on the view.',
   'my worklist just built itself.'),
 '04_network_coverage_90.png': (
   'Coverage moved to 90% — the surfaced set widens and the headline recounts.',
   'I choose how wide the net is; the screen sizes itself to the answer.',
   'more coverage when I want it, never a fixed magic number.'),
 '05_network_lens_filter.png': (
   'Lens filter narrows the map to geopolitical risk only.',
   'one dimension of risk isolated across the whole network.',
   '"where is my geopolitical exposure" answered in one click.'),
 '06_network_node_panel.png': (
   'Clicking a node opens risk factor, exposure, path from root, and the recommended action.',
   'every node carries its dollar stake.',
   'the reason and the next move travel with the alert.'),
 '07_fulltable_open.png': (
   '"See all" opens the full indented BOM — nothing is hidden, the long tail is one click away.',
   'triage without loss of audit.',
   'the complete list when I need to dig.'),
 '08_fulltable_pop_sorted.png': (
   'Sorted by POP Change — the biggest movers since last period rise to the top.',
   'the weekly staff-meeting question — better or worse than last week — answered by a column.',
   'reds on top, chase the change not the noise.'),
 '09_fulltable_lens.png': (
   'Lens filter inside the table — movement isolated by risk dimension.',
   'change, by lens, by supplier, exportable.',
   "filter to my commodity's risk and work the list."),
 '10_signals_landing.png': (
   'Signals — the daily screen: every anomaly with score, delta, exposure, break date, source.',
   'the morning view of what moved.',
   'my day starts here, not in email.'),
 '11_signals_movers.png': (
   'Sorted by biggest delta — the sharpest breaks first.',
   'attention ordered by change x dollars.',
   'worst first, always.'),
 '12_signals_drawer.png': (
   'The drawer proves it — 30-day baseline, marked breakout, and the sub-factor bar showing exactly what drove the jump.',
   'a number my board can interrogate — baseline, break, cause, source.',
   'I can defend this alert to anyone in one screen.'),
 '13_qsclive_analyzed.png': (
   "The same engine on QSC's real 792-part BOM — live data, not a demo set.",
   'this is our supply base, ranked.',
   'real parts, real suppliers, real exposure.'),
 '14_qsclive_fulltable.png': (
   'The full QSC table with POP and Lens columns — the complete network, sortable and exportable.',
   'the audit trail behind the map.',
   'everything, filterable, in one place.'),
}
ORDER = sorted(CAPTIONS.keys())

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tb.text_frame.word_wrap = True
    return tb.text_frame

# --- Slide 1: title ---
s = prs.slides.add_slide(BLANK)
bar = s.shapes.add_shape(1, 0, Inches(3.1), SW, Inches(0.08))  # MSO_SHAPE.RECTANGLE = 1
bar.fill.solid(); bar.fill.fore_color.rgb = GOLD_HEX; bar.line.fill.background()
tf = textbox(s, Inches(1), Inches(1.9), Inches(11.3), Inches(1.0))
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = 'RADAR'; r.font.size = Pt(20); r.font.bold = True; r.font.color.rgb = GOLD_HEX
tf2 = textbox(s, Inches(1), Inches(2.35), Inches(11.3), Inches(0.9))
p = tf2.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = 'RADAR — Click-by-Click Walkthrough'
r.font.size = Pt(36); r.font.bold = True; r.font.color.rgb = INK
tf3 = textbox(s, Inches(1.5), Inches(3.45), Inches(10.3), Inches(0.8))
p = tf3.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = 'From 500 suppliers to the ones that matter — and the moment one starts to slip.'
r.font.size = Pt(16); r.font.color.rgb = GREY

# --- Slides 2-15: frames ---
IMG_W = Emu(int(SW * 0.80))
IMG_H = Emu(int(IMG_W * 1080 / 1920))
IMG_L = Emu(int((SW - IMG_W) / 2))
IMG_T = Inches(0.35)

def cap_par(tf, first, label, text):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    r = p.add_run(); r.text = label; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = GOLD_HEX if label == 'WHAT CHANGED: ' else INK
    r2 = p.add_run(); r2.text = text; r2.font.size = Pt(11); r2.font.color.rgb = INK if label == 'WHAT CHANGED: ' else GREY
    p.space_after = Pt(2)

for name in ORDER:
    what, vp, scm = CAPTIONS[name]
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(os.path.join(ANNOT, name), IMG_L, IMG_T, IMG_W, IMG_H)
    tf = textbox(s, Inches(1.0), Emu(int(IMG_T + IMG_H)) + Inches(0.12), Inches(11.3), Inches(1.35))
    cap_par(tf, True, 'WHAT CHANGED: ', what)
    cap_par(tf, False, 'FOR THE VP: ', vp)
    cap_par(tf, False, 'FOR THE SUPPLY CHAIN MANAGER: ', scm)

# --- Slide 16: close ---
s = prs.slides.add_slide(BLANK)
tf = textbox(s, Inches(1.5), Inches(2.9), Inches(10.3), Inches(1.8))
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = ('RADAR does not give you more to look at. It tells you where to look — '
          'in order of money, before your supplier calls you.')
r.font.size = Pt(24); r.font.bold = True; r.font.color.rgb = INK

prs.save(OUT)
print(f'{OUT} — {len(prs.slides._sldIdLst)} slides')
