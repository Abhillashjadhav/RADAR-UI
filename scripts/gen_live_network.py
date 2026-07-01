import openpyxl, json, re
from collections import defaultdict, Counter

SRC="/root/.claude/uploads/d61e1f38-657b-593a-9a82-dbd2ebe4a45a/f32830d0-Supplier_data_file_qsc_testing_v1.xlsx"
OUT="/home/user/RADAR-UI/src/data/liveNetworkData.ts"

wb=openpyxl.load_workbook(SRC, data_only=True)
ws=wb["Sheet1"]
rows=[]
for r in range(2, ws.max_row+1):
    v=[ws.cell(row=r,column=c).value for c in range(1,14)]
    if any(x is not None for x in v): rows.append(v)

# column indices
ITEM,MPN,SUP,Z2SUP,COMM,Z2COMM,FAC,SITE,ISTAT,MSTAT,IMPACT,REV,WU = range(13)

def norm(s): return (s or '').strip()
def wu(r):
    v=r[WU]
    return float(v) if isinstance(v,(int,float)) else 1.0
def impact(r): return (norm(r[IMPACT]) or 'Medium').title()
IMPACT_W={'High':3,'Medium':2,'Low':1}
BASE={'High':78,'Medium':55,'Low':30}

def is_single_part(r):
    # MPN Sourcing Status SS = single sourced; item status all-S also single
    m=norm(r[MSTAT]).upper()
    i=norm(r[ISTAT]).upper()
    return m=='SS' or i in ('SSS',)

# commodity -> impact bucket
DELIVERY={'IC','DIODE','TRANSISTOR','CAPACITOR','RESISTOR','INDUCTOR','FERRITE','CRYSTAL','OSCILLATOR','RELAY','SWITCH'}
COST={'PCB','CONNECTOR','HARDWARE','METAL','PLASTIC','WIRE/CABLE','PACKAGING','GASKET','LABEL','ADHESIVE'}
def bucket(comm):
    c=(comm or '').upper()
    if c in DELIVERY: return 'Delivery'
    if c in COST: return 'Cost'
    return 'Compliance'

# location lookup for well-known suppliers (data has no geo)
LOC={
 'YAGEO':('New Taipei','Taiwan'),'MURATA':('Kyoto','Japan'),'SAMSUNG':('Suwon','South Korea'),
 'VISHAY':('Malvern','USA'),'VISHAY DALE':('Columbus','USA'),'PANASONIC':('Osaka','Japan'),
 'TEXAS INSTRUMENTS':('Dallas','USA'),'ON SEMICONDUCTOR':('Phoenix','USA'),'KOA':('Nagano','Japan'),
 'KEMET':('Fort Lauderdale','USA'),'BOURNS':('Riverside','USA'),'TAIYO YUDEN':('Tokyo','Japan'),
 'STACKPOLE':('Raleigh','USA'),'ROYAL OHM':('Dongguan','China'),'TDK':('Tokyo','Japan'),
 'LITTELFUSE':('Chicago','USA'),'IXYS':('Milpitas','USA'),'TE CONNECTIVITY':('Schaffhausen','Switzerland'),
 'DIODES INC.':('Plano','USA'),'ZETEX':('Oldham','UK'),'RENESAS':('Tokyo','Japan'),'NEC':('Tokyo','Japan'),
 'VTECH (DONGGUAN)':('Dongguan','China'),'VTECH (SBU6)':('Dongguan','China'),
 'GP ELECTRONICS (HUIZHOU) CO.':('Huizhou','China'),'GREENCONN':('Taoyuan','Taiwan'),
 'ABRACON':('Spicewood','USA'),'PULSE ENGINEERING':('San Diego','USA'),
 'CORNELL DUBILIER':('Liberty','USA'),'CDE/MALLORY':('Liberty','USA'),
 'NICHICON':('Kyoto','Japan'),'ROHM':('Kyoto','Japan'),'AVX':('Fountain Inn','USA'),
 'WURTH':('Niedernhall','Germany'),'MOLEX':('Lisle','USA'),'AMPHENOL':('Wallingford','USA'),
 'MICROCHIP':('Chandler','USA'),'STMICROELECTRONICS':('Geneva','Switzerland'),
 'INFINEON':('Neubiberg','Germany'),'NXP':('Eindhoven','Netherlands'),'ANALOG DEVICES':('Wilmington','USA'),
}
def loc_for(name):
    key=name.strip().upper()
    if key in LOC:
        c=LOC[key]; return {'city':c[0],'country':c[1]}
    return {'city':'—','country':'Global'}

# ---- aggregate per Supplier (T1) ----
sup_rows=defaultdict(list)
for r in rows:
    sup_rows[norm(r[SUP])].append(r)

# choke: Z2 legal entity shared by >1 distinct Supplier short-name
z2_suppliers=defaultdict(set)
z2_rows=defaultdict(list)
for r in rows:
    z2=norm(r[Z2SUP])
    if z2:
        z2_suppliers[z2].add(norm(r[SUP]))
        z2_rows[z2].append(r)
choke_z2={k for k,v in z2_suppliers.items() if len(v)>1}

def slug(s):
    return re.sub(r'[^A-Za-z0-9]+','-',s.strip()).strip('-').upper()[:40]

def agg_node(rlist, node_id, name, tier, is_choke=False, parents_count=1):
    partCount=len(rlist)
    total_wu=sum(wu(r) for r in rlist)
    raw=sum(IMPACT_W[impact(r)]*wu(r) for r in rlist)
    # risk score weighted by where-used
    tw=sum(wu(r) for r in rlist) or 1
    wscore=sum(BASE[impact(r)]*wu(r) for r in rlist)/tw
    single_share=sum(1 for r in rlist if is_single_part(r))/max(1,partCount)
    bump=10 if single_share>=0.5 else (4 if single_share>=0.25 else 0)
    score=int(max(10,min(97,round(wscore+bump))))
    is_single = single_share>=0.5
    has_high = any(impact(r)=='High' for r in rlist)
    is_spof = is_single and has_high and not is_choke
    # dominant commodity
    comm=Counter(norm(r[Z2COMM]) or norm(r[COMM]) for r in rlist).most_common(1)[0][0] or '—'
    primary=bucket(Counter(norm(r[COMM]) for r in rlist).most_common(1)[0][0])
    # lens
    if is_choke:
        lens,lensLabel=('CHK','Convergence / Sub-tier concentration')
    elif is_spof:
        lens,lensLabel=('SPF','Single-Source Continuity')
    elif is_single:
        lens,lensLabel=('SNG','Sourcing Concentration')
    else:
        lens,lensLabel=('CAP','Capacity / Demand')
    # action
    if is_choke:
        action=f"Convergence point — {parents_count} tier-1 suppliers route through {name}. Assess concentration; map alternate sub-tier capacity."
    elif is_spof:
        action=f"Single-source risk — qualify a second source for {comm}. {partCount} parts across {int(total_wu)} assemblies."
    elif is_single:
        action=f"Sourcing concentration — {int(single_share*100)}% single-sourced. Review dual-source options for {comm}."
    else:
        action=f"Monitor — {partCount} parts across {int(total_wu)} assemblies; multi-sourced."
    return {
        'id':node_id,'name':name,'tier':tier,'location':loc_for(name),
        'commodity':comm,'riskScore':score,'primaryImpact':primary,
        '_raw':raw,'_partCount':partCount,'_wu':total_wu,
        'isSPOF':is_spof,'isChokePoint':is_choke,
        'topRiskLens':lens,'topRiskLensLabel':lensLabel,'action':action,
    }

# build T1 nodes
t1_nodes={}
for sup,rlist in sup_rows.items():
    if not sup: continue
    t1_nodes[sup]=agg_node(rlist, f"T1-{slug(sup)}", sup, 1)

# build T2 choke nodes
t2_nodes={}
for z2 in choke_z2:
    t2_nodes[z2]=agg_node(z2_rows[z2], f"T2-{slug(z2)}", z2, 2, is_choke=True, parents_count=len(z2_suppliers[z2]))

# scale exposure so max T1 raw ~ 25 ($25M-equiv)
max_raw=max((n['_raw'] for n in t1_nodes.values()), default=1)
SCALE=25.0/max_raw if max_raw else 1
def exposure(n): return round(n['_raw']*SCALE,3)

# attach T2 choke under each contributing T1 (reuse same id -> choke/dedup)
# map short supplier -> set of z2 chokes
sup_to_choke=defaultdict(set)
for z2 in choke_z2:
    for sup in z2_suppliers[z2]:
        if sup in t1_nodes:
            sup_to_choke[sup].add(z2)

def emit(n, children_z2=()):
    rev=exposure(n)
    loc=n['location']
    kids=[]
    for z2 in sorted(children_z2):
        kids.append(emit(t2_nodes[z2]))
    return {
        'id':n['id'],'name':n['name'],'tier':n['tier'],
        'location':{'city':loc['city'],'country':loc['country']},
        'commodity':n['commodity'],'riskScore':n['riskScore'],
        'primaryImpact':n['primaryImpact'],
        'revenueAtRisk':rev,
        'isSPOF':n['isSPOF'],'isChokePoint':n['isChokePoint'],
        'topRiskLens':n['topRiskLens'],'topRiskLensLabel':n['topRiskLensLabel'],
        'action':n['action'],'children':kids,
    }

t1_json=[]
for sup in sorted(t1_nodes, key=lambda s:-t1_nodes[s]['_raw']):
    t1_json.append(emit(t1_nodes[sup], sup_to_choke.get(sup, ())))

root={
    'id':'QSC','name':'QSC Aerospace','tier':0,
    'location':{'city':'El Segundo','country':'USA'},
    'commodity':'OEM','riskScore':0,'primaryImpact':'Delivery',
    'revenueAtRisk':None,'isSPOF':False,'isChokePoint':False,
    'topRiskLens':'—','topRiskLensLabel':'—','action':'—','children':t1_json,
}

# stats
n_spof=sum(1 for n in t1_nodes.values() if n['isSPOF'])
n_choke=len(t2_nodes)
print(f"T1 suppliers: {len(t1_nodes)}  SPOF: {n_spof}  choke T2: {n_choke}  parts: {len(rows)}  scale: {SCALE:.4f}")

def ts(v, ind=0):
    sp='  '*ind
    if v is None: return 'null'
    if isinstance(v,bool): return 'true' if v else 'false'
    if isinstance(v,(int,float)): return repr(v)
    if isinstance(v,str): return json.dumps(v)
    if isinstance(v,list):
        if not v: return '[]'
        items=',\n'.join(sp+'  '+ts(x,ind+1) for x in v)
        return '[\n'+items+'\n'+sp+']'
    if isinstance(v,dict):
        items=',\n'.join(f'{sp}  {k}: {ts(val,ind+1)}' for k,val in v.items())
        return '{\n'+items+'\n'+sp+'}'
    raise TypeError(type(v))

header='''// AUTO-GENERATED from customer file: Supplier_data_file_qsc_testing_v1.xlsx
// 792 BOM parts -> QSC Aerospace sub-tier network. Do not edit by hand;
// regenerate with scratchpad/gen_live.py if the source file changes.
//
// DERIVATIONS (the source sheet has no revenue, geo, or lens columns):
//   revenueAtRisk  -> DERIVED EXPOSURE INDEX, not dollars.
//                     raw = sum(impactWeight[High=3,Med=2,Low=1] * WhereUsed) per node,
//                     scaled so the largest tier-1 ~= 25 (to reuse the $M-scaled UI).
//   riskScore      -> where-used-weighted base (High 78 / Med 55 / Low 30)
//                     + single-source bump.
//   isSPOF         -> majority single-sourced (MPN status SS / item SSS) AND some High-impact part.
//   isChokePoint   -> Z2 (sub-tier) legal entity shared by >1 distinct tier-1 supplier.
//   primaryImpact  -> mapped from dominant commodity family.
//   location       -> lookup for well-known suppliers; otherwise "Global".
//   topRiskLens    -> synthesized from sourcing/choke posture.
//
import type { SubTierFullNode } from './subtierMockData';

export const QSC_LIVE_NETWORK: SubTierFullNode = '''

with open(OUT,'w') as f:
    f.write(header+ts(root)+';\n')
print("wrote", OUT)
